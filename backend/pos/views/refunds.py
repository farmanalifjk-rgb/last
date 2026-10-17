from pos.models import Company, Customer, Order, OrderItem, Payment, Refund, RefundItem
from ..serializers import OrderSerializer
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from saas.permissions import HasActiveTenant, TenantRBACPermission

from rest_framework.response import Response
from decimal import Decimal
from django.db.models import Sum
from django.shortcuts import get_object_or_404
from django.utils import timezone
from django.db import transaction
from rest_framework.pagination import PageNumberPagination
import csv
from django.http import HttpResponse
from ..filters import get_filtered_refunds
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.platypus import (SimpleDocTemplate,Table,TableStyle)
from io import BytesIO
from ..pdf_helpers import add_company_header, add_report_info, add_report_title, add_summary_table, style_report_table
from ..services.inventory import update_stock
from ..helpers import sanitize_csv_cell, add_page_footer, build_export_filename



class RefundDetailView(APIView):
    permission_classes = [IsAuthenticated, HasActiveTenant, TenantRBACPermission]
    rbac_module = "pos"
    rbac_action = "view"

    def get(self, request, order_id):

        order = get_object_or_404(Order, id=order_id, tenant=request.tenant)

        items = []

        for item in OrderItem.objects.filter(order=order):

            refunded = (
                RefundItem.objects.filter(order_item=item)
                .aggregate(total=Sum("quantity"))["total"]
                or 0
            )

            items.append({

                "order_item_id": item.id,

                "product": sanitize_csv_cell(item.product.name),

                "sold_qty": item.quantity,

                "refunded_qty": refunded,

                "remaining_qty": item.quantity - refunded,

                "price": item.unit_price,

            })

        return Response({

            "order_id": order.id,

            "order_number": order.order_number,

            "customer": order.customer.name if order.customer else "Walk-in Customer",

            "items": items

        })


class RefundHistoryView(APIView):
    permission_classes = [IsAuthenticated, HasActiveTenant, TenantRBACPermission]
    rbac_module = "pos"
    rbac_action = "view"

    def get(self, request, order_id):

        refunds = Refund.objects.filter(
            tenant=request.tenant,
            order_id=order_id
        ).order_by("-created_at")

        data = []

        for refund in refunds:

            items = RefundItem.objects.filter(
                refund=refund
            )

            data.append({

                "refund_id": refund.id,

                "date": refund.created_at,

                "reason": sanitize_csv_cell(refund.reason),

                "total": refund.total_amount,

                "items":[

                    {

                        "product": sanitize_csv_cell(i.order_item.product.name),

                        "qty": i.quantity,

                        "amount": i.amount

                    }

                    for i in items

                ]

            })

        return Response(data)


class RefundReceiptView(APIView):
    permission_classes = [IsAuthenticated, HasActiveTenant, TenantRBACPermission]
    rbac_module = "pos"
    rbac_action = "view"

    def get(self, request, refund_id):

        refund = get_object_or_404(
            Refund,
            id=refund_id,
            tenant=request.tenant,
        )

        company = Company.objects.filter(tenant=request.tenant).first()

        refund_subtotal = sum(
            item.order_item.unit_price * item.quantity
            for item in refund.items.all()
        )

        if refund.order.total > 0:
            refund_ratio = refund.total_amount / refund.order.total
        else:
            refund_ratio = Decimal("0.00")

        refund_discount = refund.order.discount * refund_ratio

        refund_tax = refund.order.tax * refund_ratio   
        
        refund_total = refund_subtotal - refund_discount + refund_tax 

        return Response({

            "receipt_type": "refund",

            "company": {
                "name": company.name if company else "",
                "tagline": company.tagline if company else "",
                "address": company.address if company else "",
                "phone": company.phone if company else "",
                "email": company.email if company else "",
                "footer_message": company.footer_message if company else "",
                "website": company.website if company else "",
                "currency": company.currency if company else "Rs",
                "logo": company.logo.url if company and company.logo else None,
                "gst": company.gst if company else "",
            },

            "totals": {

                "subtotal": round(refund_subtotal,2),

                "discount": round(refund_discount,2),

                "tax": round(refund_tax,2),

                "total": round(refund_total,2),

                "paid": round(refund.total_amount,2),

                "change": 0,

            },

            "order": {
                "number": f"RE-{refund.id:06d}",
                "original_order": refund.order.order_number,
                "date": refund.created_at,
                "customer": (
                    refund.order.customer.name
                    if refund.order.customer
                    else "Walk-in Customer"
                ),
                "payment_method": refund.order.payment_method,
                "note": sanitize_csv_cell(refund.reason),
            },

            "items": [
                    {
                    "name": item.order_item.sanitize_csv_cell(product.name),
                    "qty": item.quantity,
                    "price": float(item.order_item.unit_price),
                    "subtotal": float(item.order_item.unit_price * item.quantity),
                }
                
                for item in refund.items.all()
            ]
        })   


class RefundHistoryListView(APIView):
    permission_classes = [IsAuthenticated, HasActiveTenant, TenantRBACPermission]
    rbac_module = "pos"
    rbac_action = "view"

    def get(self, request):

        refunds = get_filtered_refunds(request)          

        paginator = PageNumberPagination()
        paginator.page_size = 20

        page = paginator.paginate_queryset(refunds, request)

        data = []

        for refund in page:
        
            data.append({
            
                "id": refund.id,
                "refund_number": f"RE-{refund.id:06d}",
                "order_number": refund.order.order_number,
                "customer": (
                    refund.order.customer.name
                    if refund.order.customer
                    else "Walk-in Customer"
                ),
                "payment_method": refund.order.payment_method,
                "date": refund.created_at,
                "amount": float(refund.total_amount),
                "reason": sanitize_csv_cell(refund.reason),
            })

        return paginator.get_paginated_response(data)


class RefundDetailHistoryView(APIView):
    permission_classes = [IsAuthenticated, HasActiveTenant, TenantRBACPermission]
    rbac_module = "pos"
    rbac_action = "view"

    def get(self, request, refund_id):

        refund = get_object_or_404(
            Refund,
            id=refund_id,
            tenant=request.tenant,
        )

        return Response({

            "id": refund.id,

            "refund_number": f"RE-{refund.id:06d}",
            
            "order_id": refund.order.id,

            "order_number": refund.order.order_number,

            "customer": (
                refund.order.customer.name
                if refund.order.customer
                else "Walk-in Customer"
            ),

            "date": refund.created_at,

            "reason": sanitize_csv_cell(refund.reason),

            "total": float(refund.total_amount),

            "items": [

                {

                    "product": sanitize_csv_cell(item.order_item.product.name),

                    "qty": item.quantity,

                    "price": float(item.order_item.unit_price),

                    "amount": float(item.amount)

                }

                for item in refund.items.all()

            ]

        })
    

class RefundHistoryCSVExportView(APIView):
    permission_classes = [IsAuthenticated, HasActiveTenant, TenantRBACPermission]
    rbac_module = "pos"
    rbac_action = "view"

    def get(self, request):
        from django.conf import settings as dj_settings
        from django.http import StreamingHttpResponse
        MAX_ROWS = getattr(dj_settings, "MAX_EXPORT_ROWS", 50000)
        refunds = get_filtered_refunds(request).select_related("order", "order__customer")[:MAX_ROWS]
        def rows():
            import csv as _csv
            from io import StringIO
            buf = StringIO()
            w = _csv.writer(buf)
            w.writerow([sanitize_csv_cell(h) for h in ["Refund No", "Order No", "Customer", "Payment", "Amount", "Date"]])
            yield buf.getvalue()
            buf.seek(0); buf.truncate(0)
            for refund in refunds.iterator(chunk_size=2000):
                w.writerow([
                    sanitize_csv_cell(f"RE-{refund.id:06d}"),
                    sanitize_csv_cell(refund.order.order_number),
                    sanitize_csv_cell(refund.order.customer.name if refund.order.customer else "Walk-in Customer"),
                    sanitize_csv_cell(refund.order.payment_method),
                    sanitize_csv_cell(str(refund.total_amount)),
                    sanitize_csv_cell(refund.created_at.strftime("%Y-%m-%d %H:%M")),
                ])
                yield buf.getvalue()
                buf.seek(0); buf.truncate(0)
        filename = build_export_filename(request, "refunds", "csv")
        resp = StreamingHttpResponse(rows(), content_type="text/csv")
        resp["Content-Disposition"] = f'attachment; filename="{filename}"'
        resp["X-Export-Limit"] = str(MAX_ROWS)
        return resp
 


class RefundHistoryExportExcelView(APIView):
    permission_classes = [IsAuthenticated, HasActiveTenant, TenantRBACPermission]
    rbac_module = "pos"
    rbac_action = "view"

    def get(self, request):

        refunds = get_filtered_refunds(request)

        workbook = Workbook()
        sheet = workbook.active

        sheet.title = "Refund Orders"
        sheet.freeze_panes = "A2"

        headers = [
            "Refund No",
            "Order No",
            "Customer",
            "Payment",
            "Refund",
            "Date",
        ]

        sheet.append(headers)

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        

        for refund in refunds:

            sheet.append([
                f"RE-{refund.id:06d}",
                refund.order.order_number,
                refund.order.customer.name if refund.order.customer else "Walk-in Customer",
                refund.order.payment_method,
                refund.total_amount,
                refund.created_at.strftime("%Y-%m-%d %H:%M"),
            ])

        sheet.auto_filter.ref = sheet.dimensions

        for column_cells in sheet.columns:

            length = max(
                len(str(cell.value)) if cell.value else 0
                for cell in column_cells
            )

            sheet.column_dimensions[
                column_cells[0].column_letter
            ].width = length + 3

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        filename = build_export_filename(request,"refunds","xlsx")

        response["Content-Disposition"] = (f'attachment; filename="{filename}"')

        workbook.save(response)

        return response
 
    

class RefundHistoryExportPDFView(APIView):
    permission_classes = [IsAuthenticated, HasActiveTenant, TenantRBACPermission]
    rbac_module = "pos"
    rbac_action = "view"

    def get(self, request):

        buffer = BytesIO()

        doc = SimpleDocTemplate(buffer)

        company = Company.objects.filter(tenant=request.tenant).first()

        story = []

        add_company_header(
            story,
            company,
        )

        add_report_title(
            story,
            "REFUND HISTORY REPORT",
        ) 

        refunds = get_filtered_refunds(request)

        generated = timezone.localtime().strftime("%d-%b-%Y %I:%M %p")

        report_info = [["Generated", generated]]

        add_report_info(story,report_info,)

        table_data = [[
            "S.No",
            "Refund No",
            "Order No",
            "Customer",
            "Payment",
            "Amount",
            "Date",
        ]]

        total_refunds = 0
        refund_amount = 0

        for index, refund in enumerate(refunds, start=1):
        
            total_refunds += 1
            refund_amount += refund.total_amount

            table_data.append([
            
                index,

                f"RE-{refund.id:06d}",

                refund.order.order_number,

                refund.order.customer.name
                if refund.order.customer
                else "Walk-in Customer",

                refund.order.payment_method.title(),

                f"Rs {refund.total_amount:,.2f}",

                refund.created_at.strftime("%d-%b-%Y"),

            ])

        table = Table(
            table_data,
            colWidths=[35, 70, 80, 120, 70, 70, 80],
            repeatRows=1,
        ) 

        style_report_table(table,len(table_data),)

        table.setStyle(TableStyle([

            ("ALIGN", (0, 0), (0, -1), "CENTER"),  # S.No

            ("ALIGN", (5, 1), (5, -1), "RIGHT"),   # Amount

        ]))   

        story.append(table)  

        summary_data = [

            ["Total Refunds", str(total_refunds)],

            ["Refund Amount", f"Rs {refund_amount:,.2f}"],

        ] 

        add_summary_table(story,"REFUND SUMMARY",summary_data,)

        doc.build(
            story,
            onFirstPage=add_page_footer,
            onLaterPages=add_page_footer,
        )

        pdf = buffer.getvalue()
        buffer.close()

        response = HttpResponse(
            pdf,
            content_type="application/pdf",
        )

        filename = build_export_filename(
            request,
            "refunds",
            "pdf",
        )

        response["Content-Disposition"] = (
            f'attachment; filename="{filename}"'
        )

        return response     